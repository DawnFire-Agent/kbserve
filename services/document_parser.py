import os
import tempfile
from typing import List, TYPE_CHECKING
from docx import Document
import pdfplumber
from PIL import Image
import pytesseract
import openpyxl

if TYPE_CHECKING:
    from fastapi import UploadFile

class DocumentParser:
    def __init__(self, tesseract_path=None):
        # 从环境变量或参数获取Tesseract路径
        tesseract_path = tesseract_path or os.getenv("TESSERACT_PATH")
        if tesseract_path:
            pytesseract.pytesseract.tesseract_cmd = tesseract_path
    
    def parse_file(self, file_path: str) -> str:
        """根据文件扩展名解析不同格式的文件"""
        ext = os.path.splitext(file_path)[1].lower()
        
        if ext == '.docx':
            return self._parse_docx(file_path)
        elif ext == '.pdf':
            return self._parse_pdf(file_path)
        elif ext in ['.jpg', '.jpeg', '.png', '.bmp']:
            return self._parse_image(file_path)
        elif ext in ['.xlsx', '.xls']:
            return self._parse_excel(file_path)
        elif ext == '.txt':
            return self._parse_txt(file_path)
        else:
            raise ValueError(f"不支持的文件格式: {ext}")
    
    def _parse_docx(self, file_path: str) -> str:
        doc = Document(file_path)
        return '\n'.join([para.text for para in doc.paragraphs])
    
    def _parse_pdf(self, file_path: str) -> str:
        with pdfplumber.open(file_path) as pdf:
            return '\n'.join([page.extract_text() or '' for page in pdf.pages])
    
    def _parse_image(self, file_path: str) -> str:
        image = Image.open(file_path)
        return pytesseract.image_to_string(image)
    
    def _parse_excel(self, file_path: str) -> str:
        workbook = openpyxl.load_workbook(file_path, read_only=True)
        text = []
        for sheet_name in workbook.sheetnames:
            sheet = workbook[sheet_name]
            sheet_text = []
            for row in sheet.rows:
                row_text = [str(cell.value) for cell in row if cell.value is not None]
                if row_text:
                    sheet_text.append('\t'.join(row_text))
            if sheet_text:
                text.append(f"Sheet: {sheet_name}\n" + '\n'.join(sheet_text))
        return '\n\n'.join(text)
    
    def _parse_txt(self, file_path: str) -> str:
        with open(file_path, 'r', encoding='utf-8') as f:
            return f.read()
    
    def parse_uploaded_file(self, file: "UploadFile") -> str:
        """解析上传的文件"""
        # 读取文件内容
        file_contents = file.file.read()
        
        with tempfile.NamedTemporaryFile(delete=False, suffix=os.path.splitext(file.filename)[1]) as temp:
            temp.write(file_contents)
            temp_path = temp.name
        
        try:
            text = self.parse_file(temp_path)
        finally:
            os.unlink(temp_path)
        
        return text    